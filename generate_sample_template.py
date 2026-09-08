"""Создает пример файла template.docx с форматированием Times New Roman 14 pt."""

import docx
from docx.oxml.ns import qn
from docx.shared import Pt


def set_run_font(
    run: docx.text.run.Run, name: str = "Times New Roman", size: Pt = Pt(14), bold: bool = False
) -> None:
    run.font.name = name
    run.font.size = size
    run.font.bold = bold
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:cs"), name)


def make_template() -> None:
    doc = docx.Document()

    # Стиль Normal по умолчанию
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)

    # Заголовок
    p_h = doc.add_paragraph()
    p_h.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    r_h = p_h.add_run("АКТ ПРОВЕРКИ ИСПЫТАНИЙ № {{ ACT_NUMBER }}")
    set_run_font(r_h, size=Pt(16), bold=True)

    # Метаданные (14 pt)
    p_meta = doc.add_paragraph()
    r1 = p_meta.add_run("Дата проведения измерений: ")
    set_run_font(r1, bold=True)
    r2 = p_meta.add_run("{{ DATE }}\n")
    set_run_font(r2)

    r3 = p_meta.add_run("Наименование объекта: ")
    set_run_font(r3, bold=True)
    r4 = p_meta.add_run("{{ OBJECT_NAME }}\n")
    set_run_font(r4)

    r5 = p_meta.add_run("Ответственный оператор: ")
    set_run_font(r5, bold=True)
    r6 = p_meta.add_run("{{ OPERATOR }}\n")
    set_run_font(r6)

    p_body = doc.add_paragraph()
    r_body = p_body.add_run(
        "Настоящий акт составлен по результатам спектрального анализа "
        "сигналов на контрольных частотах C. Результаты измерений представлены ниже:"
    )
    set_run_font(r_body)

    # Метка вставки
    p_content = doc.add_paragraph()
    r_c = p_content.add_run("{{ CONTENT }}")
    set_run_font(r_c)

    # Подпись
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(24)
    r_sign = p_sign.add_run("Подпись оператора: ________________ / {{ OPERATOR }} /")
    set_run_font(r_sign)

    doc.save("template.docx")
    print("Шаблон 'template.docx' успешно обновлен (Times New Roman, 14 pt)!")


if __name__ == "__main__":
    make_template()
