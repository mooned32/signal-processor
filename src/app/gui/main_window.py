from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment
from PyQt6.QtWidgets import (
    QFileDialog,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QTableView,
    QVBoxLayout,
    QWidget,
)

from ..core.calculator import calculate_data
from .table_model import PandasTableModel


class MainWindow(QMainWindow):
    config_path: Path
    file_sn_path: Path | None
    file_n_path: Path | None
    current_df: pd.DataFrame | None
    lbl_sn: QLabel
    lbl_n: QLabel
    btn_calc: QPushButton
    table_view: QTableView
    table_model: PandasTableModel
    lbl_w: QLabel
    btn_export: QPushButton

    def __init__(self, config_path: Path) -> None:
        super().__init__()
        self.config_path = config_path
        self.file_sn_path = None
        self.file_n_path = None
        self.current_df = None

        self.setWindowTitle("Анализатор спектральных сигналов")
        self.resize(900, 620)
        self._init_ui()

    def _init_ui(self) -> None:
        main_widget = QWidget()
        layout = QVBoxLayout(main_widget)
        layout.setSpacing(10)

        # Выбор Файла 1 (Сигнал + Шум)
        sn_layout = QHBoxLayout()
        self.lbl_sn = QLabel("Файл 1 (Сигнал + Шум): Не выбран")
        btn_sn = QPushButton("Обзор...")
        _ = btn_sn.clicked.connect(self._select_file_sn)
        sn_layout.addWidget(self.lbl_sn, stretch=1)
        sn_layout.addWidget(btn_sn)
        layout.addLayout(sn_layout)

        # Выбор Файла 2 (Шум)
        n_layout = QHBoxLayout()
        self.lbl_n = QLabel("Файл 2 (Шум): Не выбран")
        btn_n = QPushButton("Обзор...")
        _ = btn_n.clicked.connect(self._select_file_n)
        n_layout.addWidget(self.lbl_n, stretch=1)
        n_layout.addWidget(btn_n)
        layout.addLayout(n_layout)

        # Кнопка Расчёта
        self.btn_calc = QPushButton("Рассчитать")
        self.btn_calc.setFixedHeight(35)
        self.btn_calc.setStyleSheet("font-weight: bold;")
        _ = self.btn_calc.clicked.connect(self._run_calculation)
        layout.addWidget(self.btn_calc)

        # Таблица результатов
        self.table_view = QTableView()
        self.table_model = PandasTableModel()
        self.table_view.setModel(self.table_model)

        header = self.table_view.horizontalHeader()
        if header is not None:
            header.setSectionResizeMode(QHeaderView.ResizeMode.Stretch)

        layout.addWidget(self.table_view)

        # Подвал
        bottom_layout = QHBoxLayout()
        self.lbl_w = QLabel("Итоговое значение W: —")
        self.lbl_w.setStyleSheet("font-size: 14px; font-weight: bold; color: #1E3A8A;")

        self.btn_export = QPushButton("Экспорт в Excel...")
        self.btn_export.setEnabled(False)
        _ = self.btn_export.clicked.connect(self._export_to_excel)

        bottom_layout.addWidget(self.lbl_w, stretch=1)
        bottom_layout.addWidget(self.btn_export)
        layout.addLayout(bottom_layout)

        self.setCentralWidget(main_widget)

    def _select_file_sn(self) -> None:
        file, _filter = QFileDialog.getOpenFileName(
            self, "Выберите файл (Сигнал + Шум)", "", "Текстовые файлы (*.txt);;Все файлы (*.*)"
        )
        if file:
            self.file_sn_path = Path(file)
            self.lbl_sn.setText(f"Файл 1: {self.file_sn_path.name}")

    def _select_file_n(self) -> None:
        file, _filter = QFileDialog.getOpenFileName(
            self, "Выберите файл (Шум)", "", "Текстовые файлы (*.txt);;Все файлы (*.*)"
        )
        if file:
            self.file_n_path = Path(file)
            self.lbl_n.setText(f"Файл 2: {self.file_n_path.name}")

    def _run_calculation(self) -> None:
        if not self.file_sn_path or not self.file_n_path:
            _ = QMessageBox.warning(
                self, "Предупреждение", "Пожалуйста, выберите оба файла перед расчетом!"
            )
            return

        try:
            df, w = calculate_data(self.file_sn_path, self.file_n_path, self.config_path)
            self.current_df = df
            self.table_model.update_data(df)

            # Объединяем ячейку W на все 20 строк таблицы
            self.table_view.clearSpans()
            col_w_idx = df.columns.get_loc("W")
            self.table_view.setSpan(0, col_w_idx, 20, 1)

            self.lbl_w.setText(f"Итоговое значение W: {w:.4f}")
            self.btn_export.setEnabled(True)
        except Exception as e:
            _ = QMessageBox.critical(self, "Ошибка расчёта", str(e))

    def _export_to_excel(self) -> None:
        if self.current_df is None:
            return
        save_path, _filter = QFileDialog.getSaveFileName(
            self, "Сохранить отчет", "report.xlsx", "Excel Files (*.xlsx)"
        )
        if not save_path:
            return

        try:
            # 1. Сохраняем исходные данные
            self.current_df.to_excel(save_path, index=False)

            # 2. Объединяем ячейки W в Excel (диапазон H2:H21)
            wb = openpyxl.load_workbook(save_path)
            ws = wb.active
            if ws is not None:
                # W — 8-й столбец (колонка H), строки со 2 по 21 (строка 1 — заголовки)
                ws.merge_cells("H2:H21")
                merged_cell = ws["H2"]
                merged_cell.alignment = Alignment(horizontal="center", vertical="center")
                wb.save(save_path)

            _ = QMessageBox.information(
                self, "Успех", f"Файл успешно сохранен с объединенной ячейкой W:\n{save_path}"
            )
        except Exception as e:
            _ = QMessageBox.critical(self, "Ошибка экспорта", str(e))
