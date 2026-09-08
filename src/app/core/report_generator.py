from pathlib import Path

import docx
import pandas as pd
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(14)
TABLE_FONT_SIZE = Pt(11)  # Для аккуратного размещения 9 столбцов в границах страницы А4


def apply_font_to_run(
    run: docx.text.run.Run,
    font_name: str = FONT_NAME,
    font_size: Pt = FONT_SIZE,
    bold: bool = False,
) -> None:
    """Применяет Times New Roman с поддержкой кириллицы к фрагменту текста."""
    run.font.name = font_name
    run.font.size = font_size
    run.font.bold = bold

    # Явная фиксация кириллического шрифта в OpenXML
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:cs"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)


def set_cell_background(cell: docx.table._Cell, hex_color: str) -> None:
    """Устанавливает цвет заливки ячейки Word."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def format_cell_text(cell: docx.table._Cell, text: str, bold: bool = False) -> None:
    """Форматирует текст внутри ячейки таблицы шрифтом Times New Roman."""
    cell.text = text
    p = cell.paragraphs[0]
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    if p.runs:
        apply_font_to_run(p.runs[0], FONT_NAME, TABLE_FONT_SIZE, bold=bold)


def insert_df_table(
    doc: docx.Document,
    df: pd.DataFrame,
    x_alerts: list[bool],
    is_extended: bool = False,
) -> docx.table.Table:
    """Вставляет стилизованную таблицу pandas DataFrame в Word документ."""
    table = doc.add_table(rows=len(df) + 1, cols=len(df.columns))
    table.style = "Table Grid"

    # 1. Заголовки
    for col_idx, col_name in enumerate(df.columns):
        cell = table.cell(0, col_idx)
        set_cell_background(cell, "E0E7FF")  # светло-синий фон шапки
        format_cell_text(cell, str(col_name), bold=True)

    # 2. Строки данных
    x_col_idx = list(df.columns).index("x") if "x" in df.columns else -1

    for row_idx in range(len(df)):
        for col_idx in range(len(df.columns)):
            cell = table.cell(row_idx + 1, col_idx)
            val = df.iloc[row_idx, col_idx]
            val_str = (
                f"{float(val):.4f}"
                if isinstance(val, (float, int)) and not isinstance(val, bool)
                else str(val)
            )

            # Подсветка x красным цветом при нарушении
            if col_idx == x_col_idx and row_idx < len(x_alerts) and x_alerts[row_idx]:
                set_cell_background(cell, "FF8A8A")

            format_cell_text(cell, val_str)

    # 3. Для расширенной таблицы объединяем столбец W на 20 строк
    if is_extended and "W" in df.columns:
        w_col_idx = list(df.columns).index("W")
        merged_cell = table.cell(1, w_col_idx).merge(table.cell(len(df), w_col_idx))
        merged_cell.text = ""
        w_val = df["W"].iloc[0]
        format_cell_text(merged_cell, f"{float(w_val):.4f}", bold=True)

    return table


def generate_docx_report(
    template_path: Path,
    output_path: Path,
    report_data: dict[str, str],
    preview_df: pd.DataFrame,
    extended_df: pd.DataFrame,
    x_alerts: list[bool],
    violating_c: list[int],
) -> None:
    """Генерирует финальный отчет Word с форматированием Times New Roman, 14."""
    if not template_path.exists():
        raise FileNotFoundError(f"Файл шаблона не найден: {template_path}")

    doc = docx.Document(template_path)

    # 0. Задаем глобальный стиль документа по умолчанию: Times New Roman, 14 pt
    style_normal = doc.styles["Normal"]
    style_normal.font.name = FONT_NAME
    style_normal.font.size = FONT_SIZE

    # 1. Подстановка текстовых меток во все параграфы
    for paragraph in doc.paragraphs:
        for key, val in report_data.items():
            tag = f"{{{{ {key} }}}}"
            tag_no_space = f"{{{{{key}}}}}"
            if tag in paragraph.text:
                paragraph.text = paragraph.text.replace(tag, val)
            if tag_no_space in paragraph.text:
                paragraph.text = paragraph.text.replace(tag_no_space, val)
        # Гарантируем Times New Roman 14 для всех фрагментов параграфа
        for run in paragraph.runs:
            apply_font_to_run(run, FONT_NAME, FONT_SIZE, bold=run.bold)

    # И в таблицы шаблона (если в них есть текстовые метки)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for key, val in report_data.items():
                    tag = f"{{{{ {key} }}}}"
                    tag_no_space = f"{{{{{key}}}}}"
                    if tag in cell.text:
                        cell.text = cell.text.replace(tag, val)
                    if tag_no_space in cell.text:
                        cell.text = cell.text.replace(tag_no_space, val)
                for p in cell.paragraphs:
                    for run in p.runs:
                        apply_font_to_run(run, FONT_NAME, FONT_SIZE, bold=run.bold)

    # 2. Поиск места вставки (по метке {{ CONTENT }} или в конец)
    target_p = None
    for p in doc.paragraphs:
        if "{{ CONTENT }}" in p.text or "{{CONTENT}}" in p.text:
            target_p = p
            p.text = ""
            break

    # 3. Вставляем таблицу превью
    p_title = doc.add_paragraph() if target_p is None else target_p.insert_paragraph_before()
    run_t = p_title.add_run("Таблица измерений (превью):")
    apply_font_to_run(run_t, FONT_NAME, FONT_SIZE, bold=True)

    _ = insert_df_table(doc, preview_df, x_alerts, is_extended=False)

    # 4. Строка статуса проверки (14 pt, Times New Roman)
    p_status = doc.add_paragraph()
    p_status.paragraph_format.space_before = Pt(14)
    p_status.paragraph_format.space_after = Pt(14)

    if not violating_c:
        run_status = p_status.add_run("При проверке не обнаружено нарушений.")
        apply_font_to_run(run_status, FONT_NAME, FONT_SIZE, bold=True)
        run_status.font.color.rgb = RGBColor(0x15, 0x80, 0x3D)  # зеленый
    else:
        c_list_str = ", ".join(str(c) for c in violating_c)
        run_status = p_status.add_run(f"При проверке обнаружены нарушения на C: {c_list_str}")
        apply_font_to_run(run_status, FONT_NAME, FONT_SIZE, bold=True)
        run_status.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)  # красный

        # 5. При наличии нарушений вставляем N расширенных таблиц
        for idx, c_val in enumerate(violating_c, 1):
            p_ext = doc.add_paragraph()
            p_ext.paragraph_format.space_before = Pt(16)
            run_ext = p_ext.add_run(f"Расширенная таблица для нарушения №{idx} (C = {c_val}):")
            apply_font_to_run(run_ext, FONT_NAME, FONT_SIZE, bold=True)

            _ = insert_df_table(doc, extended_df, x_alerts, is_extended=True)

    doc.save(output_path)
